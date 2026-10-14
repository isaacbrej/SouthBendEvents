import concurrent.futures
import glob
import json
import os
import re
from urllib.parse import urljoin, urlparse
import openpyxl
import requests
import urllib3
from bs4 import BeautifulSoup
import yaml

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)


class UniversalSelfHealingInvestigator:

    def __init__(
        self, config_path="sources.yaml", log_path="scraper_output.log"
    ):
        self.config_path = config_path
        self.log_path = log_path
        self.excel_path = self._find_master_excel()
        self.headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/122.0.0.0 Safari/537.36"
            ),
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.5",
        }

    def _find_master_excel(self):
        candidates = [
            "South_Bend_Master_Event_Sources_Cleaned.xlsx",
            "South_Bend_Master_Event_Sources_Cleaned_2.xlsx",
            "South_Bend_Master_Event_Sources_v3.xlsx",
            "South_Bend_Master_Event_Sources.xlsx",
        ]
        for candidate in candidates:
            if os.path.exists(candidate):
                return candidate
        matches = glob.glob("*South_Bend*.xlsx") + glob.glob("*Master*.xlsx")
        return matches[0] if matches else None

    def parse_log(self):
        if not os.path.exists(self.log_path):
            print(f"❌ Log file '{self.log_path}' not found.")
            return []

        with open(self.log_path, "r", encoding="utf-8") as f:
            log_content = f.read()

        problem_sources = []
        seen_names = set()

        # Build mapping of venue names to actual URLs logged in headers
        scrape_header_pattern = re.compile(
            r"🔍 Scraping: (.*?)\s*\((https?://[^\)]+)\)"
        )
        url_map = {}
        for match in scrape_header_pattern.finditer(log_content):
            name, url = match.group(1).strip(), match.group(2).strip()
            url_map[name] = url

        blocks = log_content.split("🔍 Scraping: ")

        for block in blocks[1:]:
            lines = block.strip().split("\n")
            header = lines[0]

            name_match = re.match(
                r"^(.*?)\s*\((https?://[^\)]+|[^\)]+)\)", header
            )
            if not name_match:
                continue

            venue_name = name_match.group(1).strip()
            extracted_url = name_match.group(2).strip()
            real_url = (
                extracted_url
                if extracted_url.startswith("http")
                else url_map.get(venue_name, extracted_url)
            )

            if venue_name in seen_names:
                continue

            issue_detected = None
            category = "Network / Access Error"

            if re.search(
                r"✅ Harvested 0 events from " + re.escape(venue_name), block
            ) or re.search(r"Harvested 0 events", block):
                issue_detected = "Zero Events Harvested"
                category = "Zero Events Harvested (HTTP 200 / Rendered)"
            elif "⚠️ Failed to reach" in block:
                status_match = re.search(r"Status (\d+)", block)
                status_code = (
                    status_match.group(1) if status_match else "Unknown"
                )
                issue_detected = f"Failed to Reach (Status {status_code})"
                category = "Network / Access Error"
            elif "❌ Connection failed for" in block or "Timeout" in block:
                if "NameResolutionError" in block:
                    issue_detected = "DNS / Domain Name Resolution Error"
                else:
                    issue_detected = "Connection / TLS Handshake Timeout"
                category = "Network / Access Error"
            else:
                continue

            seen_names.add(venue_name)
            source_id = re.sub(
                r"[^a-z0-9_]", "", venue_name.lower().replace(" ", "_")
            )

            problem_sources.append(
                {
                    "id": source_id,
                    "name": venue_name,
                    "url": real_url,
                    "flagged_issue": issue_detected,
                    "category": category,
                }
            )

        return problem_sources

    def discover_alternate_urls(self, base_url):
        candidates = []
        try:
            parsed = urlparse(base_url)
            root_url = f"{parsed.scheme}://{parsed.netloc}"

            res = requests.get(
                root_url, headers=self.headers, timeout=5, verify=False
            )
            if res.status_code == 200:
                soup = BeautifulSoup(res.text, "html.parser")
                for a in soup.find_all("a", href=True):
                    href = a["href"]
                    link_text = a.get_text(strip=True).lower()

                    if any(
                        kw in href.lower() or kw in link_text
                        for kw in [
                            "event",
                            "calendar",
                            "schedule",
                            "upcoming",
                            "shows",
                        ]
                    ):
                        full_alt = urljoin(root_url, href)
                        if (
                            full_alt.startswith("http")
                            and full_alt not in candidates
                            and full_alt != base_url
                        ):
                            try:
                                test_res = requests.head(
                                    full_alt,
                                    headers=self.headers,
                                    timeout=3,
                                    verify=False,
                                )
                                if test_res.status_code in [200, 301, 302]:
                                    candidates.append(full_alt)
                            except Exception:
                                pass

                        if len(candidates) >= 2:
                            break
        except Exception:
            pass

        return candidates

    def diagnose_source(self, source_info):
        url = source_info["url"]
        name = source_info["name"]
        flagged = source_info["flagged_issue"]
        category = source_info["category"]

        diag = {
            "id": source_info["id"],
            "name": name,
            "url": url,
            "flagged_issue": flagged,
            "status_code": None,
            "is_js_rendered": False,
            "has_json_ld": False,
            "replacement_url": None,
            "should_disable": False,
            "switch_to_playwright": False,
            "diagnosis": "",
            "recommended_action": "",
            "category": category,
        }

        if not url.startswith(("http://", "https://")):
            diag["diagnosis"] = (
                "URL cell contains plain text notes instead of a valid web address."
            )
            diag["recommended_action"] = (
                "Update Excel cell with valid https:// address."
            )
            return diag

        try:
            res = requests.get(
                url, headers=self.headers, timeout=5, verify=False
            )
            diag["status_code"] = res.status_code

            if res.status_code in [404, 400]:
                alts = self.discover_alternate_urls(url)
                if alts:
                    diag["replacement_url"] = alts[0]
                    diag["diagnosis"] = (
                        f"Original link returned Status {res.status_code}. Found working alternate URL: {alts[0]}"
                    )
                    diag["recommended_action"] = (
                        "Auto-healing: Replacing broken link in Excel and YAML."
                    )
                else:
                    diag["diagnosis"] = (
                        f"Link returned Status {res.status_code} and no working sub-pages were found."
                    )
                    diag["recommended_action"] = "Manual site review required."
                return diag

            if res.status_code == 403:
                diag["switch_to_playwright"] = True
                diag["diagnosis"] = (
                    "Server returned 403 Forbidden. Cloudflare/WAF anti-bot active."
                )
                diag["recommended_action"] = (
                    "Auto-healing: Setting scraping_type to 'playwright' in sources.yaml."
                )
                return diag

            html = res.text
            soup = BeautifulSoup(html, "html.parser")

            text_len = len(soup.get_text(strip=True))
            scripts_count = len(soup.find_all("script"))
            if text_len < 400 and scripts_count > 4:
                diag["is_js_rendered"] = True
                diag["switch_to_playwright"] = True

            candidate_selectors = [
                "article",
                ".event",
                ".event-card",
                ".event-item",
                ".sqs-events-collection-item",
                ".type-tribe_events",
                "a.calendarData",
                "tr.calendarRow",
            ]
            found = [s for s in candidate_selectors if soup.select(s)]

            if diag["is_js_rendered"]:
                diag["diagnosis"] = (
                    "HTML shell loaded, but content renders dynamically via JS."
                )
                diag["recommended_action"] = (
                    "Auto-healing: Enabling Playwright rendering."
                )
            elif found:
                diag["diagnosis"] = (
                    "Container elements found, but CSS selectors failed to capture text."
                )
                diag["recommended_action"] = (
                    "Refine title/date selectors in sources.yaml."
                )
            else:
                diag["diagnosis"] = (
                    "Page reachable, but uses custom HTML classes or has 0 active listings."
                )
                diag["recommended_action"] = (
                    "Inspect DOM in browser or verify off-season calendar."
                )

        except requests.exceptions.ConnectionError:
            diag["should_disable"] = True
            diag["diagnosis"] = (
                "DNS Resolution Error or domain unreachable (Closed / Defunct domain)."
            )
            diag["recommended_action"] = (
                "Auto-healing: Setting 'enabled: false' in sources.yaml."
            )
        except requests.exceptions.Timeout:
            diag["diagnosis"] = "Server connection timed out (>5s limit reached)."
            diag["recommended_action"] = (
                "Server is slow or dropping automated requests."
            )
        except Exception as e:
            diag["diagnosis"] = f"Inspection Error: {str(e)}"
            diag["recommended_action"] = "Review site manually."

        return diag

    def apply_self_healing_fixes(self, diagnostics):
        yaml_updates = 0
        excel_updates = 0
        disabled_count = 0

        yaml_data = {}
        if os.path.exists(self.config_path):
            with open(self.config_path, "r", encoding="utf-8") as f:
                yaml_data = yaml.safe_load(f) or {}

        sources_list = yaml_data.get("sources", [])
        source_map = {s["id"]: s for s in sources_list}

        wb = None
        ws = None
        if self.excel_path and os.path.exists(self.excel_path):
            try:
                wb = openpyxl.load_workbook(self.excel_path)
                ws = wb.active
            except Exception as e:
                print(f"⚠️ Could not load Excel file: {e}")

        for item in diagnostics:
            sid = item["id"]
            name = item["name"]
            replacement = item.get("replacement_url")
            disable = item.get("should_disable")
            playwright_switch = item.get("switch_to_playwright")

            if sid in source_map:
                src = source_map[sid]
                if replacement:
                    src["url"] = replacement
                    yaml_updates += 1
                if disable:
                    src["enabled"] = False
                    disabled_count += 1
                if (
                    playwright_switch
                    and src.get("scraping_type") != "playwright"
                ):
                    src["scraping_type"] = "playwright"
                    yaml_updates += 1

            if ws:
                for row in range(5, ws.max_row + 1):
                    row_name = ws.cell(row=row, column=2).value
                    if row_name and str(row_name).strip() == name:
                        if replacement:
                            ws.cell(row=row, column=3).value = replacement
                            excel_updates += 1
                        if disable:
                            ws.cell(row=row, column=6).value = (
                                "Disabled - Domain Dead"
                            )
                            excel_updates += 1

        if yaml_data and (yaml_updates > 0 or disabled_count > 0):
            with open(self.config_path, "w", encoding="utf-8") as f:
                yaml.dump(
                    yaml_data,
                    f,
                    default_flow_style=False,
                    sort_keys=False,
                    allow_unicode=True,
                )
            print(
                f"🔧 YAML Auto-Healed: Updated {yaml_updates} settings, disabled {disabled_count} dead domains."
            )

        if wb and excel_updates > 0:
            try:
                wb.save(self.excel_path)
                print(
                    f"📖 Excel Auto-Healed: Updated {excel_updates} cells in '{self.excel_path}'."
                )
            except Exception as e:
                print(f"⚠️ Failed to save Excel sheet: {e}")

    def run_investigation(self):
        targets = self.parse_log()

        if not targets:
            print("🎉 No problematic sources found in 'scraper_output.log'!")
            return

        print(
            f"🚀 Investigating {len(targets)} problematic targets concurrently...\n"
        )

        network_errors = []
        zero_event_successes = []

        with concurrent.futures.ThreadPoolExecutor(
            max_workers=10
        ) as executor:
            future_to_target = {
                executor.submit(self.diagnose_source, target): target
                for target in targets
            }
            for future in concurrent.futures.as_completed(future_to_target):
                target = future_to_target[future]
                try:
                    data = future.result()
                    if (
                        data["category"]
                        == "Zero Events Harvested (HTTP 200 / Rendered)"
                    ):
                        zero_event_successes.append(data)
                    else:
                        network_errors.append(data)

                    print(
                        f"✅ Diagnosed [{data['category']}] '{data['name']}' -> {data['diagnosis']}"
                    )
                except Exception as exc:
                    print(f"⚠️ Exception inspecting {target['name']}: {exc}")

        zero_event_successes.sort(key=lambda x: x["name"])
        network_errors.sort(key=lambda x: x["name"])

        all_diagnostics = zero_event_successes + network_errors
        self.apply_self_healing_fixes(all_diagnostics)

        categorized_report = {
            "summary": {
                "total_investigated": len(targets),
                "successful_http200_zero_events_count": len(
                    zero_event_successes
                ),
                "network_access_errors_count": len(network_errors),
            },
            "successful_http200_zero_events": zero_event_successes,
            "network_and_access_errors": network_errors,
        }

        report_file = "investigation_report.json"
        with open(report_file, "w", encoding="utf-8") as f:
            json.dump(categorized_report, f, indent=2)

        print("\n" + "=" * 70)
        print("📊 SELF-HEALED INVESTIGATION COMPLETE")
        print(
            f"   - ⚠️  HTTP 200 / Rendered with 0 Events: {len(zero_event_successes)}"
        )
        print(
            f"   - 🚫 Network / Access Errors: {len(network_errors)}"
        )
        print(f"💾 Full report saved to '{report_file}'")
        print("=" * 70)


if __name__ == "__main__":
    investigator = UniversalSelfHealingInvestigator(
        log_path="scraper_output.log"
    )
    investigator.run_investigation()