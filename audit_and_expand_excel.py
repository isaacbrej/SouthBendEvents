import concurrent.futures
import glob
import os
import re
from urllib.parse import urljoin, urlparse
import openpyxl
import requests
import urllib3
from bs4 import BeautifulSoup

# Suppress SSL warnings for local/self-signed sites
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
}

# Standard calendar/event URL paths to proactively check
CANDIDATE_PATHS = [
    "/calendar",
    "/events",
    "/event-calendar",
    "/upcoming-events",
    "/schedule",
    "/shows",
    "/upcoming",
    "/community/calendar",
    "/things-to-do",
]


def find_master_excel():
    candidates = [
        "South_Bend_Master_Event_Sources_Cleaned.xlsx",
        "South_Bend_Master_Event_Sources_Cleaned_2.xlsx",
        "South_Bend_Master_Event_Sources_v3.xlsx",
        "South_Bend_Master_Event_Sources.xlsx",
    ]
    for candidate in candidates:
        if os.path.exists(candidate):
            return candidate
    matches = glob.glob("*South_Bend*.xlsx") + glob.glob("*Master*.xlsx")
    return matches[0] if matches else None


def discover_alternate_urls(source_name, current_url):
    """Checks direct path candidates and homepage navigational links for better calendar endpoints."""
    results = {
        "name": source_name,
        "current_url": current_url,
        "discovered_alts": [],
        "status": "Checked",
    }

    if not current_url or not str(current_url).startswith(
        ("http://", "https://")
    ):
        results["status"] = "Invalid URL String"
        return results

    parsed = urlparse(current_url)
    root_url = f"{parsed.scheme}://{parsed.netloc}"
    discovered = []

    # 1. Direct path candidate checks
    for path in CANDIDATE_PATHS:
        test_url = urljoin(root_url, path)
        if test_url.rstrip("/") == current_url.rstrip("/"):
            continue

        try:
            res = requests.head(
                test_url, headers=HEADERS, timeout=4, verify=False
            )
            if res.status_code in [200, 301, 302]:
                get_res = requests.get(
                    test_url, headers=HEADERS, timeout=4, verify=False
                )
                if get_res.status_code == 200 and len(get_res.text) > 800:
                    discovered.append(test_url)
        except Exception:
            pass

        if len(discovered) >= 2:
            break

    # 2. Homepage Link Crawl (if direct paths yielded fewer than 2 candidates)
    if len(discovered) < 2:
        try:
            res = requests.get(
                root_url, headers=HEADERS, timeout=5, verify=False
            )
            if res.status_code == 200:
                soup = BeautifulSoup(res.text, "html.parser")
                for a in soup.find_all("a", href=True):
                    href = a["href"].strip()
                    link_text = a.get_text(strip=True).lower()

                    if any(
                        kw in href.lower() or kw in link_text
                        for kw in ["calendar", "event", "schedule", "shows"]
                    ):
                        full_alt = urljoin(root_url, href)
                        if (
                            full_alt.startswith("http")
                            and full_alt.rstrip("/") != current_url.rstrip("/")
                            and full_alt not in discovered
                        ):
                            discovered.append(full_alt)

                    if len(discovered) >= 2:
                        break
        except Exception:
            pass

    results["discovered_alts"] = discovered
    return results


def run_comprehensive_audit():
    excel_path = find_master_excel()
    if not excel_path:
        print("❌ No Master Excel workbook found in folder.")
        return

    print(f"📖 Opening Master Excel Workbook: '{excel_path}'")
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    # Extract all active sources from Excel
    sources_to_audit = []
    for row in range(5, ws.max_row + 1):
        name = ws.cell(row=row, column=2).value
        url = ws.cell(row=row, column=3).value

        if name and url:
            url_str = str(url).strip()
            if url_str.startswith(("http://", "https://")):
                sources_to_audit.append(
                    {"row": row, "name": str(name).strip(), "url": url_str}
                )

    print(
        f"🔍 Found {len(sources_to_audit)} sources to audit. Scanning concurrently...\n"
    )

    # Concurrently audit all sources
    audit_results = {}
    with concurrent.futures.ThreadPoolExecutor(
        max_workers=12
    ) as executor:
        future_map = {
            executor.submit(
                discover_alternate_urls, item["name"], item["url"]
            ): item
            for item in sources_to_audit
        }
        for future in concurrent.futures.as_completed(future_map):
            item = future_map[future]
            try:
                res = future.result()
                audit_results[item["row"]] = res
                alts = res["discovered_alts"]
                if alts:
                    print(
                        f"✅ [{item['name']}] Discovered alternate URL: {alts[0]}"
                    )
                else:
                    print(f"ℹ️ [{item['name']}] No new alternate URLs found.")
            except Exception as e:
                print(f"⚠️ Error auditing {item['name']}: {e}")

    print("\n✍️ Inserting new rows into Excel for discovered candidate URLs...")

    # Iterate backwards through rows to prevent row insertion index shifting
    inserted_rows_count = 0
    sorted_rows = sorted(audit_results.keys(), reverse=True)

    for row_idx in sorted_rows:
        res = audit_results[row_idx]
        alts = res.get("discovered_alts", [])

        if not alts:
            continue

        for alt_url in reversed(alts):
            insert_pos = row_idx + 1
            ws.insert_rows(insert_pos)

            # Copy all cell values from the original parent row
            for col_idx in range(1, ws.max_column + 1):
                ws.cell(row=insert_pos, column=col_idx).value = ws.cell(
                    row=row_idx, column=col_idx
                ).value

            # Update Source Name with suffix and set new Target URL in Column 3
            orig_name = ws.cell(row=row_idx, column=2).value
            ws.cell(row=insert_pos, column=2).value = (
                f"{orig_name} (Alt Endpoint)"
            )
            ws.cell(row=insert_pos, column=3).value = alt_url

            inserted_rows_count += 1

    output_excel_path = excel_path.replace(".xlsx", "_Expanded.xlsx")
    wb.save(output_excel_path)

    print("\n" + "=" * 70)
    print("🎉 COMPREHENSIVE AUDIT COMPLETE")
    print(f"  • Total sources audited: {len(sources_to_audit)}")
    print(f"  • New alternate venue rows added: {inserted_rows_count}")
    print(f"💾 Saved expanded Master Excel workbook to: '{output_excel_path}'")
    print("=" * 70)


if __name__ == "__main__":
    run_comprehensive_audit()